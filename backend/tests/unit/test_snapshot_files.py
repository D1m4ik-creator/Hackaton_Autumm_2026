import csv
import io
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from app.crud.customer_export import _WRITERS
from app.crud.snapshot_files import CSV_SNAPSHOT_HEADERS, parse_snapshot_file
import app.crud.snapshot_files as snapshot_files
from app.db.schemas.source import SnapshotIn


@pytest.fixture
def snapshot() -> SnapshotIn:
    return SnapshotIn.model_validate_json(
        (Path(__file__).resolve().parents[3] / "origin_data" / "data.json").read_bytes()
    )


def export_file(snapshot: SnapshotIn, fmt: str) -> bytes:
    rows = [SimpleNamespace(
        **account.model_dump(by_alias=False), level="low", score=10,
        expected_loss_cents=100, reason_short="Проверка",
    ) for account in snapshot.accounts]
    return _WRITERS[fmt](rows, snapshot)


def assert_same_source(actual: SnapshotIn, expected: SnapshotIn):
    for name in ("accounts", "services", "transactions"):
        assert sorted(getattr(actual, name), key=lambda item: item.id) == sorted(
            getattr(expected, name), key=lambda item: item.id
        )


@pytest.mark.parametrize("fmt", ["json", "csv", "xlsx"])
@pytest.mark.parametrize("empty", [False, True])
def test_export_import_roundtrip(snapshot: SnapshotIn, fmt: str, empty: bool):
    if empty:
        snapshot = SnapshotIn(accounts=[], services=[], transactions=[])
    assert_same_source(parse_snapshot_file(export_file(snapshot, fmt), fmt), snapshot)


@pytest.mark.parametrize("fmt", ["json", "csv", "xlsx"])
def test_roundtrip_preserves_exact_values_and_literal_formulas(snapshot: SnapshotIn, fmt: str):
    previous_id = snapshot.accounts[0].id
    snapshot.accounts[0].id = 2**63 - 1
    snapshot.accounts[0].balance = 2**63 - 1
    snapshot.accounts[0].email = '=HYPERLINK("bad");тест\n@example.com'
    for service in snapshot.services:
        if service.owner_id == previous_id:
            service.owner_id = snapshot.accounts[0].id
    snapshot.services[0].name = '=SUM(1,2); "Сервер"\nВторая строка'
    content = export_file(snapshot, fmt)
    assert_same_source(parse_snapshot_file(content, fmt), snapshot)
    if fmt == "xlsx":
        workbook = load_workbook(io.BytesIO(content))
        assert workbook["services"]["B2"].data_type == "s"
        workbook.close()


@pytest.mark.parametrize("delimiter", [";", ",", "\t"])
def test_csv_delimiters_and_bom(snapshot: SnapshotIn, delimiter: str):
    original = list(csv.reader(io.StringIO(export_file(snapshot, "csv").decode("utf-8-sig")), delimiter=";"))
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, delimiter=delimiter)
    writer.writerows(original)
    assert_same_source(parse_snapshot_file(buffer.getvalue().encode("utf-8"), "csv"), snapshot)


@pytest.mark.parametrize("fmt, content", [
    ("json", b"{broken"), ("csv", b""), ("xlsx", b"not an Excel file"),
    ("csv", "ID;Email\n1;user@example.com".encode()),
])
def test_broken_or_report_only_files_are_rejected(fmt: str, content: bytes):
    with pytest.raises(ValueError):
        parse_snapshot_file(content, fmt)


def test_report_only_xlsx_is_rejected():
    workbook = Workbook()
    workbook.active.append(["ID", "Email"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    with pytest.raises(ValueError, match="accounts, services и transactions"):
        parse_snapshot_file(buffer.getvalue(), "xlsx")


def test_csv_malformed_source_is_rejected():
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";")
    writer.writerow(CSV_SNAPSHOT_HEADERS)
    writer.writerow(["{}", "not json", "[]"])
    with pytest.raises(ValueError, match="строка 2"):
        parse_snapshot_file(buffer.getvalue().encode(), "csv")


@pytest.mark.parametrize("fmt", ["json", "csv", "xlsx"])
def test_invalid_source_fields_use_snapshot_validation(snapshot: SnapshotIn, fmt: str):
    snapshot.accounts[0].balance = -1
    with pytest.raises(ValidationError) as error:
        parse_snapshot_file(export_file(snapshot, fmt), fmt)
    assert error.value.errors()[0]["loc"] == ("accounts", 0, "balance")


def test_xlsx_formula_in_source_is_rejected(snapshot: SnapshotIn):
    workbook = load_workbook(io.BytesIO(export_file(snapshot, "xlsx")))
    workbook["accounts"]["C2"] = "=1+1"
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    with pytest.raises(ValueError, match="вместо формул"):
        parse_snapshot_file(buffer.getvalue(), "xlsx")


def test_csv_long_transaction_history(snapshot: SnapshotIn):
    data = snapshot.model_dump(mode="json")
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";")
    writer.writerow(CSV_SNAPSHOT_HEADERS)
    writer.writerow([
        json.dumps(data["accounts"][0]), json.dumps(data["services"]),
        json.dumps(data["transactions"] * 20),
    ])
    parsed = parse_snapshot_file(buffer.getvalue().encode(), "csv")
    assert len(parsed.transactions) == len(snapshot.transactions) * 20


def test_xlsx_rejects_oversized_unpacked_file(snapshot: SnapshotIn, monkeypatch):
    content = export_file(snapshot, "xlsx")
    monkeypatch.setattr(snapshot_files, "MAX_XLSX_UNPACKED_BYTES", 1)
    with pytest.raises(ValueError, match="после распаковки"):
        parse_snapshot_file(content, "xlsx")


def test_xlsx_does_not_silently_truncate_source_text(snapshot: SnapshotIn):
    snapshot.services[0].name = "a" * 32768
    with pytest.raises(HTTPException) as error:
        export_file(snapshot, "xlsx")
    assert error.value.status_code == 422
    assert "CSV или JSON" in error.value.detail
