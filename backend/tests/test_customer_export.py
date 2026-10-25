"""Выгрузка клиентов: CSV/Excel/JSON с фильтром и сортировкой как у списка."""

import csv
import io
import json
from pathlib import Path

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook

from app.crud.customer_export import HEADERS
from app.crud.snapshot_files import CSV_SNAPSHOT_HEADERS, parse_snapshot_file
from app.db.schemas.source import SnapshotIn

SNAPSHOT = json.loads(
    (Path(__file__).resolve().parents[2] / "origin_data" / "data.json").read_text(encoding="utf-8")
)
ADMIN = {"login": "admin", "password": "test_password"}


def assert_same_source(actual: SnapshotIn, expected: SnapshotIn):
    for name in ("accounts", "services", "transactions"):
        assert sorted(getattr(actual, name), key=lambda item: item.id) == sorted(
            getattr(expected, name), key=lambda item: item.id
        )


@pytest.fixture
async def loaded(client: AsyncClient) -> AsyncClient:
    assert (await client.post("/api/v1/auth/login", json=ADMIN)).status_code == 200
    assert (await client.put("/api/v1/source/snapshot", json=SNAPSHOT)).status_code == 200
    return client


async def _list_ids(client: AsyncClient, query: str = "") -> list[int]:
    page = (await client.get(f"/api/v1/customers?pageSize=200&{query}")).json()
    return [c["id"] for c in page["items"]]


async def test_export_csv_matches_list_order(loaded: AsyncClient):
    response = await loaded.get("/api/v1/customers/export?format=csv&sort=balance&order=asc")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert 'filename="customers_' in response.headers["content-disposition"]
    assert response.content.startswith("﻿".encode())

    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig")), delimiter=";"))
    assert rows[0] == HEADERS + CSV_SNAPSHOT_HEADERS
    assert [int(r[0]) for r in rows[1:]] == await _list_ids(loaded, "sort=balance&order=asc")


async def test_export_xlsx_has_header_and_all_rows(loaded: AsyncClient):
    response = await loaded.get("/api/v1/customers/export?format=xlsx")
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.content)).active
    values = list(sheet.iter_rows(values_only=True))
    assert list(values[0]) == HEADERS
    assert [r[0] for r in values[1:]] == await _list_ids(loaded)
    # Деньги в евро, а не в центах
    first = (await loaded.get(f"/api/v1/customers/{values[1][0]}")).json()
    assert values[1][5] == round(first["balance"] / 100, 2)


async def test_export_json_respects_filters(loaded: AsyncClient):
    response = await loaded.get("/api/v1/customers/export?format=json&search=0003")
    assert response.status_code == 200
    items = response.json()["items"]
    assert [i["id"] for i in items] == [3]
    assert set(items[0]) >= {"email", "level", "riskScore", "balance", "withdrawBalance", "bonusState"}


async def test_export_validation_and_auth(loaded: AsyncClient):
    assert (await loaded.get("/api/v1/customers/export?format=pdf")).status_code == 422
    await loaded.post("/api/v1/auth/logout")
    loaded.cookies.clear()
    assert (await loaded.get("/api/v1/customers/export")).status_code == 401


@pytest.mark.parametrize("fmt", ["csv", "xlsx", "json"])
async def test_export_can_be_loaded_as_snapshot(loaded: AsyncClient, fmt: str):
    response = await loaded.get(f"/api/v1/customers/export?format={fmt}&sort=id&order=asc")
    assert response.status_code == 200
    assert_same_source(parse_snapshot_file(response.content, fmt), SnapshotIn.model_validate(SNAPSHOT))
    imported = await loaded.put(
        f"/api/v1/source/snapshot/file?format={fmt}",
        content=response.content,
        headers={"Content-Type": "application/octet-stream"},
    )
    assert imported.status_code == 200
    assert {key: imported.json()[key] for key in ("accounts", "services", "transactions")} == {
        key: len(SNAPSHOT[key]) for key in ("accounts", "services", "transactions")
    }


@pytest.mark.parametrize("fmt", ["csv", "xlsx", "json"])
async def test_filtered_export_contains_only_related_source_data(loaded: AsyncClient, fmt: str):
    response = await loaded.get(f"/api/v1/customers/export?format={fmt}&search=0003")
    assert response.status_code == 200
    snapshot = parse_snapshot_file(response.content, fmt)
    expected_services = [service for service in SNAPSHOT["services"] if service["ownerId"] == 3]
    service_ids = {service["id"] for service in expected_services}
    expected = SnapshotIn.model_validate({
        "accounts": [account for account in SNAPSHOT["accounts"] if account["id"] == 3],
        "services": expected_services,
        "transactions": [tx for tx in SNAPSHOT["transactions"] if tx["serviceId"] in service_ids],
    })
    assert_same_source(snapshot, expected)


@pytest.mark.parametrize("fmt", ["csv", "xlsx", "json"])
async def test_file_merge_collision_does_not_change_data(loaded: AsyncClient, fmt: str):
    exported = await loaded.get(f"/api/v1/customers/export?format={fmt}")
    before = (await loaded.get("/api/v1/analytics/summary")).json()
    response = await loaded.put(
        f"/api/v1/source/snapshot/file?format={fmt}&mode=merge",
        content=exported.content,
        headers={"Content-Type": "application/octet-stream"},
    )
    assert response.status_code == 422
    assert (await loaded.get("/api/v1/analytics/summary")).json() == before


@pytest.mark.parametrize("fmt", ["csv", "xlsx", "json"])
async def test_file_merge_adds_snapshot_to_existing_account(loaded: AsyncClient, fmt: str):
    exported = await loaded.get(f"/api/v1/customers/export?format={fmt}")
    existing = {**SNAPSHOT["accounts"][0], "id": 101, "email": "existing@example.test"}
    assert (await loaded.put(
        "/api/v1/source/snapshot", json={"accounts": [existing], "services": [], "transactions": []}
    )).status_code == 200
    response = await loaded.put(
        f"/api/v1/source/snapshot/file?format={fmt}&mode=merge",
        content=exported.content,
        headers={"Content-Type": "application/octet-stream"},
    )
    assert response.status_code == 200
    assert response.json()["mode"] == "merge"
    assert (await loaded.get("/api/v1/analytics/summary")).json()["customersTotal"] == 11


@pytest.mark.parametrize("fmt, content", [
    ("json", b"{broken"), ("csv", b"ID;Email\n1;test@example.com"), ("xlsx", b"not xlsx"),
    ("json", b'{"accounts": [{"balance": -1}], "services": [], "transactions": []}'),
])
async def test_invalid_snapshot_file_preserves_data(loaded: AsyncClient, fmt: str, content: bytes):
    before = (await loaded.get("/api/v1/analytics/summary")).json()
    response = await loaded.put(
        f"/api/v1/source/snapshot/file?format={fmt}", content=content,
        headers={"Content-Type": "application/octet-stream"},
    )
    assert response.status_code == 422
    assert (await loaded.get("/api/v1/analytics/summary")).json() == before


async def test_snapshot_file_requires_login(loaded: AsyncClient):
    await loaded.post("/api/v1/auth/logout")
    loaded.cookies.clear()
    response = await loaded.put(
        "/api/v1/source/snapshot/file?format=csv", content=b"invalid",
        headers={"Content-Type": "application/octet-stream"},
    )
    assert response.status_code == 401
