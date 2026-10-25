"""Чтение снимков из JSON и файлов выгрузки клиентов без потери исходных данных."""

import csv
import io
import json
from typing import Literal
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from app.db.schemas.source import (
    SnapshotAccountIn,
    SnapshotIn,
    SnapshotServiceIn,
    SnapshotTransactionIn,
)

SnapshotFormat = Literal["json", "csv", "xlsx"]
SNAPSHOT_MODELS = {
    "accounts": SnapshotAccountIn,
    "services": SnapshotServiceIn,
    "transactions": SnapshotTransactionIn,
}
CSV_SNAPSHOT_HEADERS = ["snapshotAccount", "snapshotServices", "snapshotTransactions"]
# CSV может содержать историю транзакций клиента в одной ячейке.
csv.field_size_limit(max(csv.field_size_limit(), 25 * 1024 * 1024))
MAX_XLSX_UNPACKED_BYTES = 100 * 1024 * 1024


def _read_csv(content: bytes) -> dict:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("CSV должен быть сохранён в кодировке UTF-8.") from exc
    try:
        # Запятые в заголовках отчёта не должны менять разделитель файла.
        for delimiter in (";", ",", "\t"):
            reader = csv.DictReader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
            if set(CSV_SNAPSHOT_HEADERS).issubset(reader.fieldnames or []):
                break
        else:
            raise ValueError(
                "В CSV нет исходных данных снимка. Выгрузите клиентов заново: "
                "нужны столбцы snapshotAccount, snapshotServices и snapshotTransactions."
            )
        if len(set(reader.fieldnames)) != len(reader.fieldnames):
            raise ValueError("В CSV повторяются заголовки столбцов.")
        data = {name: [] for name in SNAPSHOT_MODELS}
        for number, row in enumerate(reader, start=2):
            if None in row:
                raise ValueError(f"CSV, строка {number}: число значений не совпадает с заголовком.")
            try:
                account, services, transactions = (json.loads(row[key]) for key in CSV_SNAPSHOT_HEADERS)
                if not (isinstance(account, dict) and isinstance(services, list) and isinstance(transactions, list)):
                    raise ValueError
            except (ValueError, TypeError) as exc:
                raise ValueError(f"CSV, строка {number}: повреждены исходные данные снимка.") from exc
            data["accounts"].append(account)
            data["services"].extend(services)
            data["transactions"].extend(transactions)
        return data
    except (csv.Error, IndexError) as exc:
        raise ValueError("Не удалось прочитать CSV. Проверьте формат файла и разделитель.") from exc


def _read_xlsx(content: bytes) -> dict:
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_XLSX_UNPACKED_BYTES:
                raise ValueError("Слишком большой Excel-файл после распаковки (больше 100 МиБ).")
    except BadZipFile as exc:
        raise ValueError("Не удалось прочитать Excel. Выберите файл в формате .xlsx.") from exc
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise ValueError("Не удалось прочитать Excel. Проверьте, что файл .xlsx не повреждён.") from exc
    try:
        if not set(SNAPSHOT_MODELS).issubset(workbook.sheetnames):
            raise ValueError(
                "В Excel нужны листы accounts, services и transactions с исходными данными. "
                "Для восстановления снимка выгрузите клиентов заново."
            )
        data = {}
        for name, model in SNAPSHOT_MODELS.items():
            sheet = workbook[name]
            sheet.reset_dimensions()
            rows = sheet.iter_rows(max_col=len(model.model_fields))
            headers = [cell.value for cell in next(rows, [])]
            required = {field.alias for field in model.model_fields.values() if field.is_required()}
            if not required.issubset(headers) or len(set(headers)) != len(headers):
                raise ValueError(f"Лист {name}: проверьте заголовки столбцов ({', '.join(sorted(required))}).")
            records = []
            for number, cells in enumerate(rows, start=2):
                if all(cell.value is None for cell in cells):
                    continue
                if any(cell.data_type == "f" for cell in cells):
                    raise ValueError(f"Лист {name}, строка {number}: используйте значения вместо формул.")
                # Excel хранит пустую строку как пустую ячейку; name может быть пустым.
                records.append({
                    key: ("" if key == "name" and cell.value is None else cell.value)
                    for key, cell in zip(headers, cells)
                })
            data[name] = records
        return data
    except (AttributeError, BadZipFile, KeyError, OSError, SyntaxError, TypeError) as exc:
        raise ValueError("Не удалось прочитать листы Excel. Проверьте, что файл .xlsx не повреждён.") from exc
    finally:
        workbook.close()


def parse_snapshot_file(content: bytes, fmt: SnapshotFormat) -> SnapshotIn:
    if fmt == "json":
        try:
            data = json.loads(content)
        except (ValueError, UnicodeDecodeError) as exc:
            raise ValueError("Не удалось прочитать JSON. Проверьте формат файла.") from exc
    elif fmt == "csv":
        data = _read_csv(content)
    elif fmt == "xlsx":
        data = _read_xlsx(content)
    else:
        raise ValueError("Поддерживаются только JSON, CSV и Excel (.xlsx).")
    return SnapshotIn.model_validate(data)
