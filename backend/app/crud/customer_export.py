import csv
import io
import json
from collections import defaultdict
from collections.abc import Sequence
from datetime import datetime, timezone
from typing import Literal

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils import get_column_letter
from sqlalchemy import Row
from sqlalchemy.ext.asyncio import AsyncSession

import app.dao.customer_query as customer_query
from app.crud.snapshot_files import CSV_SNAPSHOT_HEADERS, SNAPSHOT_MODELS
from app.db.schemas.source import (
    SnapshotAccountIn,
    SnapshotIn,
    SnapshotServiceIn,
    SnapshotTransactionIn,
)

ExportFormat = Literal["csv", "xlsx", "json"]

MEDIA_TYPES: dict[str, str] = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "json": "application/json",
}

LEVEL_LABELS = {"high": "Высокий", "medium": "Средний", "low": "Низкий", "no_data": "Нет данных"}

# Заголовки для людей (CSV/Excel): деньги в евро, уровень по-русски
HEADERS = [
    "ID",
    "Email",
    "Уровень риска",
    "Риск, %",
    "Ожидаемые потери, €",
    "Баланс, €",
    "К выводу, €",
    "Бонусы, €",
    "Бонусы разблокированы",
    "Всего пополнено, €",
    "Главная причина",
]

# Ячейка, начинающаяся с этих символов, в Excel станет формулой (CSV/formula injection)
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _safe_text(value: str) -> str:
    return f"'{value}" if value.startswith(_FORMULA_PREFIXES) else value


def _euro(cents: int) -> float:
    return round(cents / 100, 2)


def _table_row(row: Row) -> list:
    return [
        row.id,
        _safe_text(row.email),
        LEVEL_LABELS.get(row.level, row.level),
        row.score,
        _euro(row.expected_loss_cents),
        _euro(row.balance),
        _euro(row.withdraw_balance),
        _euro(row.bonus_balance),
        "да" if row.bonus_state == "unlocked" else "нет",
        _euro(row.total_replenished),
        _safe_text(row.reason_short),
    ]


def _json_row(row: Row) -> dict:
    """Те же ключи и единицы (центы), что в GET /customers, плюс поля счёта"""
    return {
        "id": row.id,
        "email": row.email,
        "level": row.level,
        "riskScore": row.score,
        "expectedLossCents": row.expected_loss_cents,
        "balance": row.balance,
        "withdrawBalance": row.withdraw_balance,
        "bonusBalance": row.bonus_balance,
        "bonusState": row.bonus_state,
        "totalReplenished": row.total_replenished,
        "reasonShort": row.reason_short,
    }


def _to_csv(rows: Sequence[Row], snapshot: SnapshotIn) -> bytes:
    buffer = io.StringIO()
    # «;» и BOM — чтобы Excel с русской локалью открыл файл двойным кликом без мастера импорта
    writer = csv.writer(buffer, delimiter=";")
    writer.writerow(HEADERS + CSV_SNAPSHOT_HEADERS)
    data = snapshot.model_dump(mode="json")
    accounts = {account["id"]: account for account in data["accounts"]}
    services = defaultdict(list)
    transactions = defaultdict(list)
    owners = {}
    for service in data["services"]:
        services[service["ownerId"]].append(service)
        owners[service["id"]] = service["ownerId"]
    for transaction in data["transactions"]:
        transactions[owners[transaction["serviceId"]]].append(transaction)
    for row in rows:
        source = [accounts[row.id], services[row.id], transactions[row.id]]
        writer.writerow(_table_row(row) + [json.dumps(value, ensure_ascii=False) for value in source])
    return ("﻿" + buffer.getvalue()).encode("utf-8")


def _to_xlsx(rows: Sequence[Row], snapshot: SnapshotIn) -> bytes:
    for name in SNAPSHOT_MODELS:
        for record in getattr(snapshot, name):
            if any(isinstance(value, str) and len(value) > 32767 for value in record.model_dump().values()):
                raise HTTPException(
                    status_code=422,
                    detail="Текст не помещается в ячейку Excel. Используйте CSV или JSON.",
                )
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Клиенты")
    widths = [12, 32, 14, 9, 18, 12, 12, 12, 12, 16, 48]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(_table_row(row))
    # Отдельные листы сохраняют исходные поля и единицы (центы).
    for name, model in SNAPSHOT_MODELS.items():
        source_sheet = workbook.create_sheet(name)
        headers = [field.alias for field in model.model_fields.values()]
        source_sheet.freeze_panes = "A2"
        source_sheet.append(headers)
        for record in getattr(snapshot, name):
            values = record.model_dump(mode="json")
            cells = []
            for key in headers:
                value = values[key]
                # Excel округляет большие целые: сохраняем их текстом.
                if isinstance(value, int) and abs(value) >= 10**15:
                    value = str(value)
                cell = WriteOnlyCell(source_sheet, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"
                cells.append(cell)
            source_sheet.append(cells)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _to_json(rows: Sequence[Row], snapshot: SnapshotIn) -> bytes:
    payload = {"exportedAt": datetime.now(timezone.utc).isoformat(), "items": [_json_row(row) for row in rows]}
    payload.update(snapshot.model_dump(mode="json"))
    return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")


_WRITERS = {"csv": _to_csv, "xlsx": _to_xlsx, "json": _to_json}


async def export_customers(
    session: AsyncSession,
    *,
    fmt: ExportFormat,
    levels: Sequence[str] | None,
    search: str | None,
    sort: str,
    order: str,
) -> tuple[bytes, str]:
    """Файл выгрузки клиентов под теми же фильтром и сортировкой, что таблица; возвращает (тело, имя файла)"""
    rows = await customer_query.export_customers(session, levels=levels, search=search, sort=sort, order=order)
    services, transactions = await customer_query.export_customer_sources(session, levels=levels, search=search)
    snapshot = SnapshotIn(
        accounts=[SnapshotAccountIn.model_validate(row) for row in rows],
        services=[SnapshotServiceIn.model_validate(service) for service in services],
        transactions=[SnapshotTransactionIn.model_validate(transaction) for transaction in transactions],
    )
    stamp = datetime.now(timezone.utc).strftime("%Y-%m-%d_%H-%M")
    return _WRITERS[fmt](rows, snapshot), f"customers_{stamp}.{fmt}"
